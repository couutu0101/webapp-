"""Đọc file Excel "Data khách hàng" (trang có cột "Biển số", "Ngày vào xưởng"…) → các phần SQL để nạp vào Supabase.

Chạy:  python3 cong-cu/nap_data_khach.py <file.xlsx> <thư mục ra>
Ra:    00_kiem_tra.txt (báo cáo dòng lỗi / cần xem), 01_xoa_tam.sql, 02..NN_phan_k.sql (dán dữ liệu vào app.nap_tam),
       99_nap.sql (xóa hết khách hàng – xe – hồ sơ rồi nạp lại bằng app.nap_data_khach(true)).
Đọc theo TÊN CỘT, không theo vị trí; cột lạ (vd "Cần kiểm tra") bỏ qua.
"""
import sys, os, re, json, datetime as dt
from collections import Counter
import openpyxl

COT = {  # tên cột trong file → khóa
    'biển số': 'bien_so', 'khách hàng': 'khach', 'điện thoại': 'dien_thoai', 'hãng xe': 'hang_xe', 'dòng xe': 'dong_xe',
    'màu': 'mau', 'hạn bh thân vỏ': 'han_than_vo', 'hạn bh tnds': 'han_tnds', 'mã hồ sơ': 'ma',
    'ngày vào xưởng': 'ngay_vao', 'ngày giao xe': 'ngay_giao', 'nguồn': 'nguon', 'đơn vị bảo hiểm': 'hang_bh',
    'giám định viên': 'gdv', 'tình trạng xe': 'tinh_trang', 'yêu cầu khách hàng': 'tinh_trang', 'tư vấn': 'tu_van', 'phản hồi khách': 'phan_hoi',
    'số km': 'so_km',
    'giá trị sửa chữa (đ)': 'gia_tri', 'khách còn nợ (đ)': 'con_no_khach', 'bh còn nợ (đ)': 'con_no_bh', 'ghi chú': 'ghi_chu',
}
NGAY = {'han_than_vo', 'han_tnds', 'ngay_vao', 'ngay_giao'}
SO = {'gia_tri', 'con_no_khach', 'con_no_bh', 'so_km'}


def chuan(s): return re.sub(r'\s+', ' ', str(s or '').strip().lower())


def doc_ngay(v):
    if v in (None, ''): return None
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    t = str(v).strip()
    for f in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y'):
        try: return dt.datetime.strptime(t[:10], f).date().isoformat()
        except ValueError: pass
    raise ValueError(f'ngày không đọc được "{t}"')


def doc_so(v):
    if v in (None, ''): return None
    if isinstance(v, (int, float)): return round(v)
    t = re.sub(r'[.,\s đ]', '', str(v))
    if not re.fullmatch(r'-?\d+', t): raise ValueError(f'số không đọc được "{v}"')
    return int(t)


def doc_file(duong):
    wb = openpyxl.load_workbook(duong, data_only=True)
    for ws in wb.worksheets:
        for r in range(1, 12):
            hang = [chuan(c.value) for c in ws[r]]
            if 'biển số' in hang and 'ngày vào xưởng' in hang:
                return ws, r, {i: COT[h] for i, h in enumerate(hang) if h in COT}
    raise SystemExit('Không thấy trang nào có cột "Biển số" và "Ngày vào xưởng".')


def main(duong, ra):
    os.makedirs(ra, exist_ok=True)
    ws, r_dau, ban_do = doc_file(duong)
    dong, loi, can_xem = [], [], []
    for r in range(r_dau + 1, ws.max_row + 1):
        gia_tri = {k: ws.cell(r, i + 1).value for i, k in ban_do.items()}
        if all(v in (None, '') for v in gia_tri.values()): continue
        if str(ws.cell(r, 2).value or '').strip().lower().startswith('tổng') or str(ws.cell(r, 1).value or '').strip().lower().startswith('tổng'):
            continue
        d = {}
        try:
            for k, v in gia_tri.items():
                if k in NGAY: d[k] = doc_ngay(v)
                elif k in SO: d[k] = doc_so(v)
                else:
                    v = None if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip()
                    d[k] = v or None
        except ValueError as e:
            loi.append(f'Dòng {r}: {e}'); continue
        if not d.get('bien_so'): loi.append(f'Dòng {r}: thiếu biển số – bỏ'); continue
        dt_ = re.sub(r'\D', '', d.get('dien_thoai') or '')
        if len(dt_) == 9 and not dt_.startswith('0'): dt_ = '0' + dt_          # Excel mất số 0 đầu
        if dt_ and not re.fullmatch(r'0\d{9}', dt_): can_xem.append(f'Dòng {r}: SĐT "{d["dien_thoai"]}" không đủ 10 số – để trống')
        d['dien_thoai'] = dt_ if re.fullmatch(r'0\d{9}', dt_ or '') else None
        n = chuan(d.get('nguon'))
        d['nguon'] = 'khach' if ('khách' in n or 'tiền mặt' in n or n == 'tm') else ('bao_hiem' if n or d.get('hang_bh') else 'khach')   # "Tiền mặt" = khách tự trả
        for k in ('con_no_khach', 'con_no_bh'):
            if (d.get(k) or 0) < 0: can_xem.append(f'Dòng {r}: {k} âm ({d[k]}) – để 0'); d[k] = 0
        if (d.get('con_no_khach') or d.get('con_no_bh')) and not d.get('ngay_vao'):
            loi.append(f'Dòng {r}: có nợ nhưng thiếu ngày vào xưởng – bỏ'); continue
        d['_dong'] = r
        dong.append(d)
    trung = Counter((re.sub(r'[^0-9A-Z]', '', d['bien_so'].upper()), d.get('ngay_vao')) for d in dong if d.get('ngay_vao'))
    for (bs, ng), c in trung.items():
        if c > 1: can_xem.append(f'Xe {bs} có {c} dòng cùng ngày vào {ng} – sẽ thành {c} hồ sơ')
    # Thông tin xe – khách lặp ở nhiều dòng của cùng một xe: nếu các dòng ghi khác nhau thì lấy dòng mới nhất, báo để xem
    TEN_COT = {'khach': 'Khách hàng', 'dien_thoai': 'Điện thoại', 'hang_xe': 'Hãng xe', 'dong_xe': 'Dòng xe', 'mau': 'Màu',
               'han_than_vo': 'Hạn BH thân vỏ', 'han_tnds': 'Hạn BH TNDS'}
    theo_xe = {}
    for d in dong: theo_xe.setdefault(re.sub(r'[^0-9A-Z]', '', d['bien_so'].upper()), []).append(d)
    for bs, ds in theo_xe.items():
        ds = sorted(ds, key=lambda d: d.get('ngay_vao') or '')
        for k, ten in TEN_COT.items():
            gt = [d.get(k) for d in ds if d.get(k)]
            if len(set(gt)) > 1:
                can_xem.append(f'Xe {d["bien_so"] if (d := ds[-1]) else bs}: cột "{ten}" ghi khác nhau giữa các dòng ({" / ".join(dict.fromkeys(map(str, gt)))}) – lấy dòng mới nhất "{gt[-1]}"')
    so_xe = len({re.sub(r'[^0-9A-Z]', '', d['bien_so'].upper()) for d in dong})
    tong_nk = sum(d.get('con_no_khach') or 0 for d in dong); tong_nb = sum(d.get('con_no_bh') or 0 for d in dong)
    with open(os.path.join(ra, '00_kiem_tra.txt'), 'w', encoding='utf-8') as f:
        f.write(f'File: {os.path.basename(duong)} – trang "{ws.title}"\n')
        f.write(f'{len(dong)} dòng hợp lệ, {so_xe} xe, {sum(1 for d in dong if d.get("ngay_vao"))} lần vào xưởng\n')
        vnd = lambda n: f'{n:,.0f}'.replace(',', '.')
        f.write(f'Nợ cũ: khách {vnd(tong_nk)} đ, bảo hiểm {vnd(tong_nb)} đ\n')
        f.write(f'\nDÒNG BỊ BỎ ({len(loi)}):\n' + '\n'.join(loi) + f'\n\nCẦN XEM ({len(can_xem)}):\n' + '\n'.join(can_xem) + '\n')
    for x in dong: x.pop('_dong')
    open(os.path.join(ra, '01_xoa_tam.sql'), 'w').write("delete from app.nap_tam where loai = 'khach_hang';\nselect 'xoa tam xong' as ket_qua;\n")
    phan, cur, co = [], [], 0
    for x in dong:
        j = json.dumps(x, ensure_ascii=False, separators=(',', ':'))
        if cur and co + len(j) > 11000: phan.append(cur); cur, co = [], 0
        cur.append(j); co += len(j) + 1
    if cur: phan.append(cur)
    for k, p in enumerate(phan, 1):
        open(os.path.join(ra, f'{k + 1:02d}_phan_{k}.sql'), 'w', encoding='utf-8').write(
            f"insert into app.nap_tam (loai, phan, dong) values ('khach_hang', {k}, $j$[{','.join(p)}]$j$::jsonb);\n"
            f"select 'phan {k} xong | dong=' || (select sum(jsonb_array_length(dong)) from app.nap_tam where loai = 'khach_hang') as ket_qua;\n")
    open(os.path.join(ra, '99_nap.sql'), 'w').write(
        f"-- Kỳ vọng: {len(dong)} dòng, {so_xe} xe, nợ khách {tong_nk}, nợ BH {tong_nb}\n"
        "select 'nap_xong | ' || app.nap_data_khach(true)::text as ket_qua;\n")
    print(open(os.path.join(ra, '00_kiem_tra.txt'), encoding='utf-8').read())
    print(f'{len(phan)} phần SQL trong {ra}')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
